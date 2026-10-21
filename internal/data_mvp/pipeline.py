"""Synthetic-only local prototype. No server, authentication, AI or network access.

Never use with customer data. Tenant checks are data-contract tests, not production security.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import sqlite3
import uuid
import zipfile
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path

FIELDS = ['row_key', 'site_id', 'metric_id', 'definition_version', 'observed_on', 'value']
MAX_BYTES = 2_000_000
MAX_ROWS = 10_000


class DataError(ValueError):
    """Non-sensitive error code; no input values are echoed."""


def now():
    return datetime.now(timezone.utc).isoformat()


def identifier(value):
    if not isinstance(value, str) or not re.fullmatch(r'[A-Za-z0-9_-]{1,80}', value):
        raise DataError('INVALID_ID')
    return value


def iso_date(value):
    try:
        value = str(value)
        if date.fromisoformat(value).isoformat() != value:
            raise ValueError
        return value
    except (ValueError, TypeError):
        raise DataError('INVALID_DATE') from None


def read_rows(path):
    path = Path(path)
    if path.stat().st_size > MAX_BYTES:
        raise DataError('FILE_TOO_LARGE')
    raw = path.read_bytes()
    if path.suffix.lower() == '.csv':
        try:
            rows = list(csv.reader(io.StringIO(raw.decode('utf-8-sig')), strict=True))
        except (UnicodeError, csv.Error):
            raise DataError('CSV_ENCODING_OR_FORMAT') from None
    elif path.suffix.lower() == '.xlsx':
        try:
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                names = archive.namelist()
                if any('vbaproject' in n.lower() or 'externallinks/' in n.lower() for n in names):
                    raise DataError('ACTIVE_WORKBOOK_REJECTED')
                if sum(z.file_size for z in archive.infolist()) > 20_000_000:
                    raise DataError('WORKBOOK_EXPANSION_LIMIT')
                if any(z.flag_bits & 1 for z in archive.infolist()):
                    raise DataError('ENCRYPTED_WORKBOOK_REJECTED')
                for name in names:
                    if name.lower().endswith('.rels') and re.search(br'TargetMode\s*=\s*[\x22\x27]External[\x22\x27]', archive.read(name), re.I):
                        raise DataError('EXTERNAL_RELATIONSHIP_REJECTED')
            from openpyxl import load_workbook
            book = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
            try:
                if len(book.worksheets) != 1:
                    raise DataError('ONE_SHEET_REQUIRED')
                rows = []
                for row in book.active.iter_rows():
                    if any(cell.data_type == 'f' for cell in row):
                        raise DataError('FORMULA_REJECTED')
                    rows.append([cell.value if cell.value is not None else '' for cell in row])
                    if len(rows) > MAX_ROWS + 1:
                        raise DataError('ROW_LIMIT')
            finally:
                book.close()
        except ImportError:
            raise DataError('XLSX_DEPENDENCY_MISSING') from None
        except Exception as exc:
            if isinstance(exc, DataError):
                raise
            raise DataError('INVALID_WORKBOOK') from None
    else:
        raise DataError('UNSUPPORTED_FILE_TYPE')
    if not rows or rows[0] != FIELDS:
        raise DataError('HEADER_MISMATCH')
    if not 1 <= len(rows) - 1 <= MAX_ROWS:
        raise DataError('ROW_LIMIT')
    if any(len(row) != len(FIELDS) for row in rows[1:]):
        raise DataError('COLUMN_COUNT')
    return raw, [dict(zip(FIELDS, row)) for row in rows[1:]]


class Store:
    def __init__(self, root, client_id):
        if not re.fullmatch(r'SYN-[A-Za-z0-9_-]{1,60}', client_id):
            raise DataError('SYNTHETIC_ONLY')
        self.client_id = client_id
        directory = Path(root) / client_id
        directory.mkdir(parents=True, exist_ok=True)
        self.db = sqlite3.connect(directory / 'synthetic.sqlite3')
        self.db.row_factory = sqlite3.Row
        self.db.executescript(Path(__file__).with_name('schema.sql').read_text())
        tenants = [r[0] for r in self.db.execute('SELECT client_id FROM tenant')]
        if tenants and tenants != [client_id]:
            self.db.close()
            raise DataError('TENANT_MISMATCH')
        self.db.execute('INSERT OR IGNORE INTO tenant VALUES (?)', (client_id,))
        self.db.commit()

    def authorize(self, client_id):
        if client_id != self.client_id:
            raise DataError('TENANT_MISMATCH')

    def add_metric(self, client_id, site_id, metric_id, version, unit, definition):
        self.authorize(client_id)
        for item in (site_id, metric_id, version):
            identifier(item)
        if not unit.strip() or not definition.strip():
            raise DataError('METRIC_DEFINITION_REQUIRED')
        with self.db:
            self.db.execute('INSERT OR IGNORE INTO sites VALUES (?,?)', (site_id, client_id))
            self.db.execute('INSERT INTO metrics VALUES (?,?,?,?,?)',
                            (site_id, metric_id, version, unit, definition))

    def ingest(self, client_id, source_id, path, origin, period_start, period_end):
        self.authorize(client_id)
        identifier(source_id)
        start, end = iso_date(period_start), iso_date(period_end)
        if start > end or not origin.strip():
            raise DataError('SOURCE_METADATA_REQUIRED')
        raw, rows = read_rows(path)
        digest = hashlib.sha256(raw).hexdigest()
        seen_keys, seen_points, clean = set(), set(), []
        for number, row in enumerate(rows, 2):
            try:
                for field in FIELDS[:4]:
                    row[field] = identifier(str(row[field]))
                observed = iso_date(row['observed_on'])
                if not start <= observed <= end:
                    raise DataError('OUTSIDE_PERIOD')
                value = Decimal(str(row['value']))
                if not value.is_finite() or len(str(row['value'])) > 80:
                    raise DataError('INVALID_VALUE')
                point = (row['site_id'], row['metric_id'], row['definition_version'], observed)
                if row['row_key'] in seen_keys or point in seen_points:
                    raise DataError('DUPLICATE_ROW')
                seen_keys.add(row['row_key']); seen_points.add(point)
                if not self.db.execute('SELECT 1 FROM metrics WHERE site_id=? AND metric_id=? AND definition_version=?', point[:3]).fetchone():
                    raise DataError('UNKNOWN_METRIC_OR_SITE')
                clean.append((row['row_key'], *point, str(value)))
            except (DataError, InvalidOperation) as exc:
                code = str(exc) if isinstance(exc, DataError) else 'INVALID_VALUE'
                raise DataError(f'{code}:ROW_{number}') from None
        with self.db:
            latest = self.db.execute('SELECT version,sha256,period_start,period_end,origin FROM sources WHERE source_id=? ORDER BY version DESC LIMIT 1', (source_id,)).fetchone()
            if latest and latest['sha256'] == digest and (latest['period_start'], latest['period_end'], latest['origin']) == (start, end, origin):
                return {'source_id': source_id, 'version': latest['version'], 'duplicate': True}
            version = latest['version'] + 1 if latest else 1
            self.db.execute('INSERT INTO sources VALUES (?,?,?,?,?,?,?,?,?,?)',
                            (source_id, version, digest, now(), origin, start, end, 'pending', None, None))
            self.db.executemany('INSERT INTO observations VALUES (?,?,?,?,?,?,?,?)',
                                [(source_id, version, *row) for row in clean])
            self.db.execute('INSERT INTO audit(occurred_at,action,reference_id,actor_role) VALUES (?,?,?,?)',
                            (now(), 'ingest_pending', f'{source_id}:v{version}', 'test_operator'))
        return {'source_id': source_id, 'version': version, 'duplicate': False, 'rows': len(clean), 'state': 'pending'}

    def review(self, client_id, source_id, version, reviewer_role, approve=True):
        self.authorize(client_id)
        identifier(reviewer_role)
        with self.db:
            changed = self.db.execute('UPDATE sources SET state=?,reviewer=?,reviewed_at=? WHERE source_id=? AND version=? AND state=?',
                ('approved' if approve else 'rejected', reviewer_role, now(), source_id, version, 'pending')).rowcount
            if changed != 1:
                raise DataError('PENDING_SOURCE_REQUIRED')
            self.db.execute('INSERT INTO audit(occurred_at,action,reference_id,actor_role) VALUES (?,?,?,?)',
                (now(), 'approve' if approve else 'reject', f'{source_id}:v{version}', reviewer_role))

    def report(self, client_id, source_id):
        self.authorize(client_id)
        source = self.db.execute('SELECT * FROM sources WHERE source_id=? ORDER BY version DESC LIMIT 1', (source_id,)).fetchone()
        if not source or source['state'] != 'approved':
            raise DataError('LATEST_SOURCE_NOT_APPROVED')
        records = self.db.execute('SELECT o.*,m.unit,m.definition FROM observations o JOIN metrics m USING(site_id,metric_id,definition_version) WHERE o.source_id=? AND o.version=? ORDER BY o.observed_on,o.row_key',
                                 (source_id, source['version'])).fetchall()
        report_id = 'SYN-R-' + uuid.uuid4().hex
        with self.db:
            self.db.execute('INSERT INTO report_runs VALUES (?,?,?,?,?)',
                            (report_id, source_id, source['version'], now(), 'draft'))
        return {'synthetic_only': True, 'client_id': client_id, 'report_id': report_id,
                'state': 'draft', 'source': dict(source), 'records': [dict(row) for row in records],
                'limits': 'No automatic aggregation, causality, customer delivery or AI analysis.'}

    def close(self):
        self.db.close()
